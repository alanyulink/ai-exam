#!/usr/bin/env python3
"""生成今日答案修改记录 Word 文档（含完整题干、选项、原答案和改后答案）"""

import json, re

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

# ---------- 读取题库数据 ----------
DATA_JS = '/Users/alanyu/学习/ai考试/js/data.js'
with open(DATA_JS, 'r', encoding='utf-8') as f:
    content = f.read()

match = re.search(r'const QUESTION_DATA = ({.*?});', content, re.DOTALL)
data = json.loads(match.group(1))

q_map = {q['id']: q for q in data['questions']}

# 题目 ID → (原答案, 改后答案)
changes = {
    114: ('D', 'B'),
    118: ('A', 'C'),
    124: ('B', 'C'),
    143: ('B', 'A'),
    179: ('A', 'C'),
    195: ('C', 'D'),
    200: ('ABCD', 'ABC'),
    256: ('B', 'A'),
    258: ('B', 'A'),
    269: ('B', 'A'),
    277: ('B', 'A'),
}

target_ids = list(changes.keys())

type_map = {'single': '单选题', 'multi': '多选题', 'judge': '判断题'}

# ---------- 生成 Word ----------
doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = '微软雅黑'
font.size = Pt(11)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

# 标题
title = doc.add_heading('题库答案修改记录', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('修改日期：2026-07-25  |  共 11 题（单选 6 题，多选 1 题，判断 4 题）')
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

doc.add_paragraph()

for idx, qid in enumerate(target_ids, 1):
    q = q_map[qid]
    qtype = type_map.get(q['type'], q['type'])
    old_ans, new_ans = changes[qid]

    # ---- 题头 ----
    p = doc.add_paragraph()
    run = p.add_run(f'【{idx:02d}】 题号 {qid}（{qtype}）')
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x00, 0x52, 0xCC)

    # ---- 答案变更 ----
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(f'原答案：{old_ans}  →  改后答案：{new_ans}')
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    # ---- 题干 ----
    p = doc.add_paragraph()
    run = p.add_run(f'题目：{q["content"]}')
    run.font.size = Pt(11)

    # ---- 选项 ----
    for opt in q['options']:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.5)
        run = p.add_run(f"{opt['label']}. {opt['content']}")
        run.font.size = Pt(11)
        # 正确答案加粗标绿
        if opt['label'] in new_ans:
            run.bold = True
            run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)

    doc.add_paragraph()  # 间隔

# 保存
output_path = '/Users/alanyu/学习/ai考试/题库答案修改记录_20260725.docx'
doc.save(output_path)
print(f'✅ 文档已生成：{output_path}')